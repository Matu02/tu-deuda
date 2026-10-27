from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from exports import excel_bytes


def test_excel_preserves_all_rows_numbers_and_literal_text():
    frame = pd.DataFrame({
        "CUIT": ["00123456789"] * 501,
        "razón_social": ["=Empresa"] * 501,
        "monto": [1234.56] * 501,
        "fechaPago": [None] * 501,
    })
    sheet = load_workbook(BytesIO(excel_bytes(frame))).active
    assert sheet.max_row == 502
    assert [cell.value for cell in sheet[1]] == list(frame.columns)
    assert sheet["A2"].value == "00123456789"
    assert sheet["B2"].value == "=Empresa"
    assert sheet["B2"].data_type == "s"
    assert sheet["C502"].value == 1234.56
    assert sheet["C502"].data_type == "n"
    assert sheet["D2"].value is None


def test_empty_excel_keeps_headers():
    sheet = load_workbook(BytesIO(excel_bytes(pd.DataFrame(columns=["CUIT", "estado"])))).active
    assert list(sheet.values) == [("CUIT", "estado")]
