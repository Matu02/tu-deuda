from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from comparison import (INCOMPARABLE, NEW, POSSIBLE_PAYMENT, UNCHANGED,
                        compare_checks, read_previous)
from exports import excel_bytes


CUIT = "30000000007"


def check(number=1, **kwargs):
    return {"CUIT": CUIT, "razon_social": "Empresa", "entidad": 11,
            "nroCheque": number, "fechaRechazo": "2026-01-01", "monto": 100,
            "fechaPago": None, **kwargs}


def snapshot(state="ok", **kwargs):
    return {"fecha": "2026-01-20", "cuits": [CUIT], "fecha_rechazo": None,
            "resultados": {CUIT: {"estados_consulta": {"cheques": state}}}, **kwargs}


def compare(before, after, snap=None):
    return compare_checks(pd.DataFrame(before), pd.DataFrame(after), snap or snapshot(), "2026-01-19")


def test_individual_checks_and_partial_company_payments():
    detail, summary = compare([check(1), check(2)], [check(2), check(3)])
    assert detail["resultado"].tolist() == [POSSIBLE_PAYMENT, UNCHANGED, NEW]
    assert summary.iloc[0]["cantidad_posibles_pagos"] == 1
    assert summary.iloc[0]["monto_posibles_pagos"] == 100
    assert detail.iloc[0]["antiguedad_dias"] == 19
    assert detail.iloc[0]["anterior.fechaRechazo"] == "2026-01-01"


def test_all_disappear_and_previously_paid_is_excluded_from_totals():
    detail, summary = compare([check(1), check(2, fechaPago="2026-01-15")], [])
    assert detail["resultado"].tolist() == [POSSIBLE_PAYMENT] * 2
    assert detail["nuevo_posible_pago"].tolist() == [True, False]
    assert summary.iloc[0]["cantidad_posibles_pagos"] == 1


@pytest.mark.parametrize("state", ["error", "parcial", "pendiente"])
def test_incomplete_checks_never_report_payment(state):
    detail, summary = compare([check()], [], snapshot(state))
    assert detail.iloc[0]["resultado"] == INCOMPARABLE
    assert summary.empty


def test_debt_error_does_not_block_complete_checks():
    snap = snapshot()
    snap["resultados"][CUIT].update(estado="parcial", estados_consulta={"cheques": "ok", "actual": "error"})
    detail, _ = compare([check()], [], snap)
    assert detail.iloc[0]["resultado"] == POSSIBLE_PAYMENT


@pytest.mark.parametrize("snap", [snapshot(cuits=[]), snapshot(fecha_rechazo="2026-01-02")])
def test_out_of_scope_checks_are_incomparable(snap):
    detail, summary = compare([check()], [], snap)
    assert detail.iloc[0]["resultado"] == INCOMPARABLE
    assert summary.empty


def test_identity_normalization_ignores_amount_name_and_payment_changes():
    old = check(CUIT="30-00000000-7", entidad="011.0", nroCheque="0001", fechaRechazo="2026-01-01T00:00:00")
    new = check(monto=300, razon_social="Otro nombre", fechaPago="2026-01-20")
    detail, summary = compare([old], [new])
    assert detail["resultado"].tolist() == [UNCHANGED]
    assert detail.iloc[0]["monto_actual"] == 300
    assert summary.empty


@pytest.mark.parametrize("before,after", [
    ([check(), check()], [check()]),
    ([check()], [check(), check()]),
    ([check(nroCheque=None)], []),
    ([check(monto="inválido")], []),
    ([check()], [check(nroCheque=None)]),
])
def test_ambiguous_or_invalid_records_never_report_payment(before, after):
    detail, summary = compare(before, after)
    assert set(detail["resultado"]) == {INCOMPARABLE}
    assert summary.empty


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_import_existing_exports(extension):
    frame = pd.DataFrame([check()])
    content = excel_bytes(frame) if extension == "xlsx" else frame.to_csv(index=False).encode("utf-8-sig")
    imported = read_previous(content, f"cheques.{extension}")
    detail, _ = compare_checks(imported, frame, snapshot(), "2026-01-19")
    assert detail["resultado"].tolist() == [UNCHANGED]


def test_empty_import_with_headers_and_empty_comparison():
    frame = pd.DataFrame(columns=list(check()))
    imported = read_previous(frame.to_csv(index=False).encode(), "cheques.csv")
    detail, summary = compare_checks(imported, frame, snapshot(), "2026-01-19")
    assert detail.empty and summary.empty


def test_invalid_file_and_date():
    with pytest.raises(ValueError, match="Faltan columnas"):
        read_previous(b"CUIT,monto\n30000000007,100", "resumen.csv")
    with pytest.raises(ValueError, match="previa"):
        compare_checks(pd.DataFrame(), pd.DataFrame(), snapshot(), "2026-01-20")


def test_export_includes_more_than_preview_limit_and_literal_text():
    detail, _ = compare([check(i, razon_social="=Empresa") for i in range(501)], [])
    sheet = load_workbook(BytesIO(excel_bytes(detail))).active
    assert sheet.max_row == 502
    column = list(detail.columns).index("razon_social") + 1
    assert sheet.cell(502, column).value == "=Empresa"
    assert sheet.cell(502, column).data_type == "s"


@pytest.mark.parametrize("rejection,age", [("2026-01-05", 15), ("2026-01-04", 16)])
def test_fifteen_days_are_context_only(rejection, age):
    detail, _ = compare([check(fechaRechazo=rejection)], [])
    assert detail.iloc[0]["resultado"] == POSSIBLE_PAYMENT
    assert detail.iloc[0]["antiguedad_dias"] == age
